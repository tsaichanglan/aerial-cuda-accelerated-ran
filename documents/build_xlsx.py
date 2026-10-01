import re, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))  # documents/
# Checkout of aerial-cuda-accelerated-ran whose source is analyzed.
# Override with: python build_xlsx.py <path>  or  set AERIAL_RAN_ROOT=<path>
ROOT = (sys.argv[1] if len(sys.argv) > 1
        else os.environ.get("AERIAL_RAN_ROOT", r"D:\aerial-cuda-accelerated-ran"))
if not os.path.isfile(os.path.join(ROOT, "cuPHY", "src", "cuphy", "cuphy_api.h")):
    sys.exit(f"aerial-cuda-accelerated-ran checkout not found at {ROOT}; pass its path as the first argument")
OUT = os.path.join(HERE, "cuPHY_5G_NR_Channels_Parameters.xlsx")

F = {
    "H":  "cuPHY/src/cuphy/cuphy.h",
    "A":  "cuPHY/src/cuphy/cuphy_api.h",
    "PR": "cuPHY/src/cuphy_channels/prach_rx.cpp",
    "PRK":"cuPHY/src/cuphy/prach_receiver/prach_receiver.hpp",
    "PT": "cuPHY/src/cuphy_channels/pdsch_tx.cpp",
    "PU": "cuPHY/src/cuphy_channels/pusch_rx.cpp",
    "CT": "cuPHY/src/cuphy_channels/csirs_tx.cpp",
    "SRS":"cuPHY/src/cuphy/srs_chEst/srs_chEst.hpp",
    "CE": "cuPHY/src/cuphy/ch_est/ch_est.cu",
    "SC": "cuPHY-CP/gt_common_libs/slot_command/include/slot_command/slot_command.hpp",
}
_cache = {}
def lines(k):
    if k not in _cache:
        with open(os.path.join(ROOT, F[k]), encoding="utf-8", errors="replace") as fh:
            _cache[k] = fh.read().splitlines()
    return _cache[k]

def ref(src):
    """src = 'KEY:symbol' or 'KEY@anchor:symbol' -> 'path:line'"""
    if not src:
        return ""
    key, sym = src.split(":", 1)
    anchor = None
    if "@" in key:
        key, anchor = key.split("@", 1)
    L = lines(key)
    start = 0
    if anchor:
        for i, l in enumerate(L):
            if anchor in l:
                start = i
                break
    pat = re.compile(r"(?<![A-Za-z0-9_])" + re.escape(sym) + r"(?![A-Za-z0-9_])")
    for i in range(start, len(L)):
        if pat.search(L[i]):
            return f"{F[key]}:{i+1}"
    print("WARN: not found", src, file=sys.stderr)
    return F[key]

# Basis labels
CT_ = "Compile-time limit"
EN = "Enforced (validation)"
DOC = "API doc comment"
SPEC = "3GPP/FAPI range (informational)"
NOTE64 = "Build with ENABLE_64C raises this value."

# Row: (channel/signal, group, parameter, cuPHY field/macro, allowed values / range, max, basis, src, notes)
DL = []
def dl(*a): DL.append(a)
UL = []
def ul(*a): UL.append(a)

# ============================ DOWNLINK ============================
# --- PDSCH ---
P = "PDSCH"
dl(P,"Cell group","Cells per PDSCH object (cell group)","PDSCH_MAX_CELLS_PER_CELL_GROUP","1..64","64",EN,"H:PDSCH_MAX_CELLS_PER_CELL_GROUP","Also bounded per slot by nMaxCellsPerSlot.")
dl(P,"Cell group","UEs (TBs) per cell group","PDSCH_MAX_UES_PER_CELL_GROUP","1..192","192 (256 with 64C)",EN,"H:PDSCH_MAX_UES_PER_CELL_GROUP",NOTE64 + " One CW per UE assumed.")
dl(P,"Cell group","UEs per cell","PDSCH_MAX_UES_PER_CELL","1..64","64",CT_,"H:PDSCH_MAX_UES_PER_CELL","")
dl(P,"Cell group","UE groups per cell group","PDSCH_MAX_UE_GROUPS_PER_CELL_GROUP","1..128","128 (256 with 64C)",EN,"H:PDSCH_MAX_UE_GROUPS_PER_CELL_GROUP",NOTE64)
dl(P,"Cell group","Heterogeneous LDPC encoder configs per cell group","PDSCH_MAX_HET_LDPC_CONFIGS_SUPPORTED","1..64","64",EN,"H:PDSCH_MAX_HET_LDPC_CONFIGS_SUPPORTED","")
dl(P,"Static","Max PRBs in DL BWP","cuphyPdschStatPrms_t.nMaxPrb","1..273 (0 = default 273)","273",EN,"A@_cuphyPdschStatPrms:nMaxPrb","")
dl(P,"Static","Max CBs per TB","cuphyPdschStatPrms_t.nMaxCBsPerTB / MAX_N_CBS_PER_TB_SUPPORTED","1..152 (0 = default)","152",EN,"H:MAX_N_CBS_PER_TB_SUPPORTED","")
dl(P,"Static","TB CRC source","read_TB_CRC","true / false","-",DOC,"A:read_TB_CRC","true: CRC read from input instead of computed.")
dl(P,"Static","Pipeline mode","pipeline_processing_mode","Full slot / AAS / post-FEC","-",DOC,"A:pipeline_processing_mode","Only full processing is supported by cuPHY-CP.")
dl(P,"Cell dynamic","Slot number","cuphyPdschCellDynPrm_t.slotNum","0..319","319",DOC,"A@_cuphyPdschCellDynPrm:slotNum","")
dl(P,"Cell dynamic","Test model mode","testModel","0 / 1","1",DOC,"A@_cuphyPdschCellDynPrm:testModel","1 = PN23 payload.")
dl(P,"Time allocation","PDSCH start symbol","pdschStartSym","0..13","13",DOC,"A@_cuphyPdschCellDynPrm:pdschStartSym","")
dl(P,"Time allocation","PDSCH symbols (DMRS + data)","nPdschSym","1..14","14",DOC,"A@_cuphyPdschCellDynPrm:nPdschSym","")
dl(P,"Time allocation","DMRS symbol bitmask","dmrsSymLocBmsk","14-bit mask","0x3FFF",DOC,"A@_cuphyPdschCellDynPrm:dmrsSymLocBmsk","")
dl(P,"Freq allocation","Resource allocation type","resourceAlloc","0 (bitmap) / 1 (contiguous)","1",DOC,"A@_cuphyPdschUeGrpPrm:resourceAlloc","")
dl(P,"Freq allocation","RB bitmap size","rbBitmap / MAX_RBMASK_BYTE_SIZE","36 bytes (273 RBs)","36 B",CT_,"H:MAX_RBMASK_BYTE_SIZE","")
dl(P,"Freq allocation","Start PRB","startPrb","0..274","274",DOC,"A@_cuphyPdschUeGrpPrm:startPrb","Not valid for RA type 0.")
dl(P,"Freq allocation","Number of PRBs","nPrb","1..275 (practical max 273)","273",DOC,"A@_cuphyPdschUeGrpPrm:nPrb","Bounded by nMaxPrb <= 273.")
dl(P,"DMRS","DMRS type","dmrsType","Type 1 only","-",DOC,"A@_cuphyPdschDmrsPrm:dmrsType","")
dl(P,"DMRS","CDM groups without data","nDmrsCdmGrpsNoData","1..3","3",DOC,"A@_cuphyPdschDmrsPrm:nDmrsCdmGrpsNoData","")
dl(P,"DMRS","DMRS symbols per slot","MAX_N_DMRSSYMS_SUPPORTED","1..4","4",CT_,"H:MAX_N_DMRSSYMS_SUPPORTED","")
dl(P,"DMRS","DMRS scrambling ID","dmrsScrmId","0..65535","65535",DOC,"A@_cuphyPdschUePrm:dmrsScrmId","")
dl(P,"DMRS","SCID","scid","0 / 1","1",DOC,"A@_cuphyPdschUePrm:scid","")
dl(P,"DMRS","DMRS port bitmask","dmrsPortBmsk","Bitmask; port 0 = LSB","-",DOC,"A@_cuphyPdschUePrm:dmrsPortBmsk","nlAbove16 extends above 16 layers.")
dl(P,"DMRS","DMRS reference point","refPoint","0 / 1","1",DOC,"A@_cuphyPdschUePrm:refPoint","")
dl(P,"UE","Layers per UE","nUeLayers","1..8","8",DOC,"A@_cuphyPdschUePrm:nUeLayers","")
dl(P,"UE","Layers per TB (codeword)","MAX_DL_LAYERS_PER_TB","1..4","4",CT_,"H:MAX_DL_LAYERS_PER_TB","")
dl(P,"UE","RNTI","rnti","1..65535","65535",DOC,"A@_cuphyPdschUePrm:rnti","")
dl(P,"UE","Data scrambling ID","dataScramId","0..65535","65535",DOC,"A@_cuphyPdschUePrm:dataScramId","")
dl(P,"UE","Codewords per UE","nCw","1..2","2",DOC,"A@_cuphyPdschUePrm:nCw","")
dl(P,"UE","Power scaling","beta_dmrs / beta_qam","float","-",DOC,"A@_cuphyPdschUePrm:beta_dmrs","Fronthaul amplitude scaling.")
dl(P,"Codeword","MCS table index","mcsTableIndex","0..2 (38.214 Tables 5.1.3.1-1/2/3)","2",DOC,"A@_cuphyPdschCwPrm:mcsTableIndex","Only used for optional TB-size check.")
dl(P,"Codeword","MCS index","mcsIndex","0..31","31",DOC,"A@_cuphyPdschCwPrm:mcsIndex","")
dl(P,"Codeword","Modulation order Qm","qamModOrder","2 (QPSK), 4 (16QAM), 6 (64QAM), 8 (256QAM)","8",DOC,"A@_cuphyPdschCwPrm:qamModOrder","")
dl(P,"Codeword","Target code rate","targetCodeRate","codeRate*1024*10","-",DOC,"A@_cuphyPdschCwPrm:targetCodeRate","Enables MCS > 28.")
dl(P,"Codeword","Redundancy version","rv","0..3","3",DOC,"A@_cuphyPdschCwPrm:rv","")
dl(P,"Codeword","LBRM PRBs","n_PRB_LBRM","{32,66,107,135,162,217,273}","273",DOC,"A@_cuphyPdschCwPrm:n_PRB_LBRM","")
dl(P,"Codeword","LBRM max layers","maxLayers","1..4","4",DOC,"A@_cuphyPdschCwPrm:maxLayers","")
dl(P,"Codeword","LBRM max Qm","maxQm","6 / 8","8",DOC,"A@_cuphyPdschCwPrm:maxQm","")
dl(P,"Transport block","Max TB size","MAX_BYTES_PER_TRANSPORT_BLOCK","<= 159773 bytes","159773 B",CT_,"H:MAX_BYTES_PER_TRANSPORT_BLOCK","Sized for 4 layers, 273 PRB, MCS 27, 1 DMRS symbol, +24 CRC.")
dl(P,"Transport block","Rate-matched bits per CB (Er)","PDSCH_MAX_ER_PER_CB_BITS","<= 256000","256000",CT_,"H:PDSCH_MAX_ER_PER_CB_BITS","Large Er only in adaptive retransmission.")
dl(P,"Precoding","Precoder ports","MAX_DL_PORTS / cuphyPmW_t.nPorts","1..32","32",CT_,"H:MAX_DL_PORTS","Precoder matrix is layers x ports (fp16 complex).")
dl(P,"Rate matching","CSI-RS params co-scheduled (RE rate-match)","nCsiRsPrms / CUPHY_CSIRS_MAX_NUM_PARAMS","<= 32 per cell","32 per cell",CT_,"H:CUPHY_CSIRS_MAX_NUM_PARAMS","")

# --- PDCCH ---
P = "PDCCH"
dl(P,"Cell group","CORESETs per cell","CUPHY_PDCCH_N_MAX_CORESETS_PER_CELL","1..40","40",CT_,"H:CUPHY_PDCCH_N_MAX_CORESETS_PER_CELL","Max CORESETs per slot = nMaxCellsPerSlot x 40.")
dl(P,"CORESET","DCIs per CORESET","nDci / CUPHY_PDCCH_MAX_DCIS_PER_CORESET","1..91","91",CT_,"H:CUPHY_PDCCH_MAX_DCIS_PER_CORESET","")
dl(P,"CORESET","CORESET symbols","n_sym","1..3","3",DOC,"A@_cuphyPdcchCoresetDynPrm:n_sym","")
dl(P,"CORESET","Start symbol","start_sym","0..13","13",SPEC,"A@_cuphyPdcchCoresetDynPrm:start_sym","")
dl(P,"CORESET","Start RB","start_rb","0..272","272",SPEC,"A@_cuphyPdcchCoresetDynPrm:start_rb","")
dl(P,"CORESET","Frequency-domain resources","freq_domain_resource","45-bit bitmap (6-RB units)","64-bit field",SPEC,"A@_cuphyPdcchCoresetDynPrm:freq_domain_resource","")
dl(P,"CORESET","CCE-to-REG mapping","interleaved","0 (non-interleaved) / 1 (interleaved)","1",DOC,"A@_cuphyPdcchCoresetDynPrm:interleaved","")
dl(P,"CORESET","REG bundle size","bundle_size","2, 3, 6 (REGs)","6",SPEC,"A@_cuphyPdcchCoresetDynPrm:bundle_size","")
dl(P,"CORESET","Interleaver size","interleaver_size","2, 3, 6","6",DOC,"A@_cuphyPdcchCoresetDynPrm:interleaver_size","")
dl(P,"CORESET","Shift index","shift_index","0..274","274",SPEC,"A@_cuphyPdcchCoresetDynPrm:shift_index","")
dl(P,"CORESET","CORESET type","coreset_type","0 / 1","1",DOC,"A@_cuphyPdcchCoresetDynPrm:coreset_type","")
dl(P,"DCI","Aggregation level","aggr_level","1, 2, 4, 8, 16","16",CT_,"H:CUPHY_PDCCH_MAX_AGGREGATION_LEVEL","5 aggregation levels.")
dl(P,"DCI","DCI payload bits (A)","Npayload","12..140","140",CT_,"H:CUPHY_PDCCH_POLAR_A_MAX","Min from CUPHY_PDCCH_POLAR_A_MIN.")
dl(P,"DCI","Polar K (payload + CRC)","CUPHY_PDCCH_POLAR_K_MAX","36..164","164",CT_,"H:CUPHY_PDCCH_POLAR_K_MAX","")
dl(P,"DCI","CRC length","CUPHY_PDCCH_N_CRC_BITS","24","24",CT_,"H:CUPHY_PDCCH_N_CRC_BITS","")
dl(P,"DCI","DCI payload buffer","CUPHY_PDCCH_MAX_DCI_PAYLOAD_BYTES","<= 20 bytes","20 B",CT_,"H:CUPHY_PDCCH_MAX_DCI_PAYLOAD_BYTES","")
dl(P,"DCI","Tx bits per DCI","CUPHY_PDCCH_MAX_TX_BITS_PER_DCI","2 x 9 x 6 x AL","1728",CT_,"H:CUPHY_PDCCH_MAX_TX_BITS_PER_DCI","At AL = 16.")
dl(P,"DCI","RNTI (CRC / scrambling)","rntiCrc / rntiBits","0..65535","65535",SPEC,"A:rntiCrc","")
dl(P,"DCI","DMRS scrambling ID","dmrs_id","0..65535","65535",SPEC,"A:dmrs_id","")
dl(P,"DCI","Power scaling","beta_qam / beta_dmrs","float","-",DOC,"A:beta_qam","")
dl(P,"DCI","Precoding","enablePrcdBf / pmwPrmIdx","0 / 1; single-layer precoder up to 32 ports","32 ports",CT_,"H:MAX_DL_PORTS","")

# --- SSB ---
P = "SSB (PSS/SSS/PBCH)"
dl(P,"Cell","SSBs per cell per slot","CUPHY_SSB_MAX_SSBS_PER_CELL_PER_SLOT","1..3","3",CT_,"H:CUPHY_SSB_MAX_SSBS_PER_CELL_PER_SLOT","")
dl(P,"Cell","Physical cell ID","NID","0..1007","1007",SPEC,"A@_cuphyPerCellSsbDynPrms:NID","")
dl(P,"Cell","Half-frame index","nHF","0 / 1","1",DOC,"A@_cuphyPerCellSsbDynPrms:nHF","")
dl(P,"Cell","L_max","Lmax","4, 8, 64","64",DOC,"A@_cuphyPerCellSsbDynPrms:Lmax","")
dl(P,"Cell","SFN","SFN","0..1023","1023",SPEC,"A@_cuphyPerCellSsbDynPrms:SFN","")
dl(P,"Cell","k_SSB subcarrier offset","k_SSB","0..31","31",DOC,"A@_cuphyPerCellSsbDynPrms:k_SSB","")
dl(P,"Cell","Subcarriers per slot","nF","0..273*12-1","3275",DOC,"A@_cuphyPerCellSsbDynPrms:nF","")
dl(P,"SS block","Start subcarrier f0","f0","[0, 273*12-240)","3035",DOC,"A@_cuphyPerSsBlockDynPrms:f0","PSS/SSS start at f0 + 56.")
dl(P,"SS block","Start symbol t0","t0","0..10","10",DOC,"A@_cuphyPerSsBlockDynPrms:t0","SSB spans 4 symbols.")
dl(P,"SS block","Block index","blockIndex","0..Lmax-1","63",DOC,"A@_cuphyPerSsBlockDynPrms:blockIndex","")
dl(P,"SS block","PSS/SSS/PBCH scaling","beta_pss / beta_sss","float","-",DOC,"A@_cuphyPerSsBlockDynPrms:beta_pss","")
dl(P,"SS block","SSB size","CUPHY_SSB_NF x CUPHY_SSB_NT","240 subcarriers x 4 symbols","-",CT_,"H:CUPHY_SSB_NF","")
dl(P,"PBCH","MIB bits","CUPHY_SSB_N_MIB_BITS","24","24",CT_,"H:CUPHY_SSB_N_MIB_BITS","")
dl(P,"PBCH","PBCH payload bits","CUPHY_SSB_N_PBCH_PAYLOAD_BITS","32 (56 with CRC)","32",CT_,"H:CUPHY_SSB_N_PBCH_PAYLOAD_BITS","")
dl(P,"PBCH","Polar-encoded bits","CUPHY_SSB_N_PBCH_POLAR_ENCODED_BITS","512","512",CT_,"H:CUPHY_SSB_N_PBCH_POLAR_ENCODED_BITS","")
dl(P,"PBCH","Rate-matched / scrambled bits","CUPHY_SSB_N_PBCH_SCRAMBLING_SEQ_BITS","864","864",CT_,"H:CUPHY_SSB_N_PBCH_SCRAMBLING_SEQ_BITS","")
dl(P,"PSS/SSS","Sequence length","CUPHY_SSB_N_SS_SEQ_BITS","127","127",CT_,"H:CUPHY_SSB_N_SS_SEQ_BITS","")
dl(P,"PBCH DMRS","Gold sequence length","CUPHY_SSB_N_DMRS_SEQ_BITS","288","288",CT_,"H:CUPHY_SSB_N_DMRS_SEQ_BITS","")

# --- CSI-RS ---
P = "CSI-RS (Tx)"
dl(P,"Cell","Antenna ports","nTxAnt / CUPHY_CSIRS_MAX_ANTENNA_PORTS","1..32","32",EN,"CT:CUPHY_CSIRS_MAX_ANTENNA_PORTS","")
dl(P,"Cell","DL BWP PRBs","nPrbDlBwp","1..273","273",EN,"CT:nPrbDlBwp","")
dl(P,"Cell","CSI-RS resources per cell","nRrcParams / CUPHY_CSIRS_MAX_NUM_PARAMS","1..32","32",CT_,"H:CUPHY_CSIRS_MAX_NUM_PARAMS","")
dl(P,"Resource","Start RB","startRb","0..272","272",EN,"CT:startRb","startRb + nRb <= 273.")
dl(P,"Resource","Number of RBs","nRb","1..273-startRb","273",EN,"CT:nRb","")
dl(P,"Resource","Table 7.4.1.5.3-1 row","row","1..18","18",EN,"H:CUPHY_CSIRS_SYMBOL_LOCATION_TABLE_LENGTH","")
dl(P,"Resource","First symbol L0 / L1","symbL0 / symbL1","0..13","13",DOC,"A@_cuphyCsirsRrcDynPrm:symbL0","")
dl(P,"Resource","Frequency density","freqDensity","0: 0.5 (even RB), 1: 0.5 (odd RB), 2: 1, 3: 3","3",DOC,"A@_cuphyCsirsRrcDynPrm:freqDensity","")
dl(P,"Resource","Frequency-domain allocation bitmap","freqDomain","Bitmap (LSB first)","-",DOC,"A@_cuphyCsirsRrcDynPrm:freqDomain","")
dl(P,"Resource","Scrambling ID","scrambId","0..1023","1023",SPEC,"A@_cuphyCsirsRrcDynPrm:scrambId","")
dl(P,"Resource","CSI type","csiType","0: TRS, 1: NZP CSI-RS, 2: ZP CSI-RS","2",DOC,"H:NZP_CSI_RS","Comment says only NZP is currently supported.")
dl(P,"Resource","CDM type","cdmType","0: noCDM, 1: FD-CDM2, 2: CDM4 (FD2-TD2), 3: CDM8 (FD2-TD4)","3",DOC,"H:CDM8_FD2_TD4","")
dl(P,"Resource","Power scaling","beta","float","-",DOC,"A@_cuphyCsirsRrcDynPrm:beta","")
dl(P,"Resource","Precoding","enablePrcdBf / pmwPrmIdx","0 / 1","<= 32 x cells matrices",EN,"CT:nPrecodingMatrices","")

# --- BFW (DL) ---
P = "BFW (DL/UL beamforming weights)"
dl(P,"Static","Regularization lambda","lambda","float > 0","-",DOC,"A@_cuphyBfwStatPrms:lambda","Regularized zero-forcing (RZF).")
dl(P,"Static","gNB antennas","nMaxGnbAnt","1..64","64",CT_,"H:MAX_N_ANTENNAS_SUPPORTED","uint8 field.")
dl(P,"Static","PRB groups","nMaxPrbGrps / CUPHY_BFW_N_MAX_PRB_GRPS","1..137","137",CT_,"H:CUPHY_BFW_N_MAX_PRB_GRPS","(273+1)/min PRB group size 2.")
dl(P,"Static","UE groups per pipeline","nMaxUeGrps / CUPHY_BFW_COEF_COMP_N_MAX_USER_GRPS","1..72","72",CT_,"H:CUPHY_BFW_COEF_COMP_N_MAX_USER_GRPS","24 UE groups/cell x 3 cells.")
dl(P,"Static","Heterogeneous configs","CUPHY_BFW_COEF_COMP_N_MAX_HET_CFGS","1..8","8",CT_,"H:CUPHY_BFW_COEF_COMP_N_MAX_HET_CFGS","")
dl(P,"Static","Compression bit width","compressBitwidth","0 (none) / 9 (9-bit BFP)","9",DOC,"A@_cuphyBfwStatPrms:compressBitwidth","Other values not supported.")
dl(P,"Static","Compression scaling","beta","float","-",DOC,"A@_cuphyBfwStatPrms:beta","")
dl(P,"Static","Power normalization algorithm","bfwPowerNormAlg_selector","Selector","-",DOC,"A@_cuphyBfwStatPrms:bfwPowerNormAlg_selector","")
dl(P,"UE group","Layers per UE group","nBfLayers / BFW_COEF_COMP_N_MAX_LAYERS_PER_USER_GRP","1..32","32",CT_,"SC:BFW_COEF_COMP_N_MAX_LAYERS_PER_USER_GRP","Now defined in cuPHY-CP.")
dl(P,"UE group","PRB group size","bfwPrbGrpSize",">= 2","-",CT_,"H:CUPHY_BFW_MIN_PRB_GRP_SIZE","")

# ============================ UPLINK ============================
P = "PUSCH"
ul(P,"Cell group","Cells per slot","nMaxCellsPerSlot / MAX_CELLS_PER_SLOT","1..20","20 (40 with 64C)",CT_,"H:MAX_CELLS_PER_SLOT",NOTE64)
ul(P,"Cell group","TBs per cell group","nMaxTbs / MAX_N_TBS_PER_CELL_GROUP_SUPPORTED","1..192","192 (256 with 64C)",EN,"PU:m_maxNTbs",NOTE64)
ul(P,"Cell group","UE groups per cell group","nUeGrps / MAX_N_USER_GROUPS_SUPPORTED","1..192","192 (256 with 64C)",EN,"PU:MAX_N_USER_GROUPS_SUPPORTED",NOTE64)
ul(P,"Static","Rx antennas","nMaxRx / MAX_N_ANTENNAS_SUPPORTED","1..64","64",EN,"PU:m_maxNRx","")
ul(P,"Static","PRBs","nMaxPrb / MAX_N_PRBS_SUPPORTED","1..273","273",EN,"PU:m_maxNPrbAlloc","")
ul(P,"Static","CBs per TB","nMaxCbsPerTb / MAX_N_CBS_PER_TB_SUPPORTED","1..152","152",EN,"PU:m_maxNCbsPerTb","")
ul(P,"Static","BBU layers (all UEs)","MAX_N_BBU_LAYERS_PUSCH_SUPPORTED","1..16","16",CT_,"H:MAX_N_BBU_LAYERS_PUSCH_SUPPORTED","")
ul(P,"Static","Channel estimation algorithm","chEstAlgo","0: legacy MMSE, 1: multi-stage MMSE + delay est, 2: RKHS, 3: LS only","3",DOC,"H:PUSCH_CH_EST_ALGO_TYPE_LS_ONLY","See the Algorithms sheet.")
ul(P,"Static","Equalizer algorithm","eqCoeffAlgo","0: RZF, 1: noise-diag MMSE, 2: MMSE-IRC, 3: MMSE-IRC shrink RBLW, 4: MMSE-IRC shrink OAS","4",DOC,"H:PUSCH_EQ_ALGO_TYPE_MMSE_IRC_SHRINK_OAS","")
ul(P,"Static","LDPC max-iteration algorithm","ldpcMaxNumItrAlgo","0: fixed, 1: LUT, 2: per UE","2",DOC,"H:LDPC_MAX_NUM_ITR_ALGO_TYPE_PER_UE","")
ul(P,"Static","LDPC decoder mode","useCbLdpcDecoder","0: TB decoder, 1: CB decoder","1",DOC,"A:useCbLdpcDecoder","")
ul(P,"Static","Polar decoder list size (UCI)","polarDcdrListSz","1..8 (PUSCH default 1)","8",CT_,"H:CUPHY_POLAR_DECODER_LIST_SIZE","")
ul(P,"Static","Feature flags","enableCfoCorrection, enableToEstimation, enablePuschTdi, enableDftSOfdm, enableRssiMeasurement, enableSinrMeasurement, enableEarlyHarq, enableMassiveMIMO, ...","0 / 1 each","-",DOC,"A:enableCfoCorrection","")
ul(P,"Static","LDPC heterogeneous configs","nMaxLdpcHetConfigs","Configurable","-",DOC,"A:nMaxLdpcHetConfigs","")
ul(P,"UE group","Layers per UE group","CUPHY_PUSCH_RX_MAX_N_LAYERS_PER_UE_GROUP","1..8","8",CT_,"H:CUPHY_PUSCH_RX_MAX_N_LAYERS_PER_UE_GROUP","")
ul(P,"UE group","UEs per UE group (MU-MIMO)","CUPHY_PUSCH_RX_MAX_N_UE_PER_UE_GROUP","1..8","8",CT_,"H:CUPHY_PUSCH_RX_MAX_N_UE_PER_UE_GROUP","1 layer each.")
ul(P,"UE group","Start PRB / nPrb","startPrb / nPrb","0..272 / 1..273","273",EN,"A@_cuphyPuschUeGrpPrm:startPrb","")
ul(P,"UE group","Start symbol","puschStartSym","0..13","13",SPEC,"A@_cuphyPuschUeGrpPrm:puschStartSym","")
ul(P,"UE group","Symbols (DMRS + data)","nPuschSym","1..14","14",SPEC,"A@_cuphyPuschUeGrpPrm:nPuschSym","")
ul(P,"UE group","DMRS symbol bitmask","dmrsSymLocBmsk","14-bit mask","0x3FFF",DOC,"A@_cuphyPuschUeGrpPrm:dmrsSymLocBmsk","")
ul(P,"UE group","RSSI symbol bitmask","rssiSymLocBmsk","14-bit mask; 0 disables","0x3FFF",DOC,"A@_cuphyPuschUeGrpPrm:rssiSymLocBmsk","")
ul(P,"DMRS","DMRS type","dmrsType","Type A mapping only","-",DOC,"A@_cuphyPuschDmrsPrm:dmrsType","")
ul(P,"DMRS","Additional positions","dmrsAddlnPos","0..3","3",SPEC,"A@_cuphyPuschDmrsPrm:dmrsAddlnPos","")
ul(P,"DMRS","Max length","dmrsMaxLen","1 / 2","2",SPEC,"A@_cuphyPuschDmrsPrm:dmrsMaxLen","")
ul(P,"DMRS","DMRS symbols per slot","N_MAX_DMRS_SYMS","1..4","4",CT_,"H:N_MAX_DMRS_SYMS","")
ul(P,"DMRS","Channel estimates in time","CUPHY_PUSCH_RX_MAX_N_TIME_CH_EST","1..4","4",CT_,"H:CUPHY_PUSCH_RX_MAX_N_TIME_CH_EST","")
ul(P,"DMRS","CDM groups without data","nDmrsCdmGrpsNoData","1..3","3",SPEC,"A@_cuphyPuschDmrsPrm:nDmrsCdmGrpsNoData","")
ul(P,"DMRS","DMRS scrambling ID","dmrsScrmId","0..65535","65535",SPEC,"A@_cuphyPuschDmrsPrm:dmrsScrmId","")
ul(P,"DMRS","SCID","scid","0 / 1","1",SPEC,"A@_cuphyPuschUePrm:scid","")
ul(P,"UE","PDU bitmap","pduBitmap","b0 data, b1 UCI, b2 PTRS, b3 DFT-s, b4 SCH, b5 CSI-P2","-",DOC,"A@_cuphyPuschUePrm:pduBitmap","")
ul(P,"UE","Transform precoding (DFT-s-OFDM)","enableTfPrcd","0 / 1","1",DOC,"A@_cuphyPuschUePrm:enableTfPrcd","")
ul(P,"UE","PUSCH identity","puschIdentity","0..1007","1007",SPEC,"A@_cuphyPuschUePrm:puschIdentity","")
ul(P,"UE","Group/sequence hopping","groupOrSequenceHopping","0 / 1 / 2","2",SPEC,"A@_cuphyPuschUePrm:groupOrSequenceHopping","")
ul(P,"UE","Low-PAPR group / sequence number","lowPaprGroupNumber / lowPaprSequenceNumber","0..29 / 0..1","29 / 1",SPEC,"A@_cuphyPuschUePrm:lowPaprGroupNumber","")
ul(P,"UE","MCS table index","mcsTableIndex","0..4 (Tables 5.1.3.1-1/2/3, 6.1.4.1-1/2)","4",DOC,"A@_cuphyPuschUePrm:mcsTableIndex","Only used for TB-size check.")
ul(P,"UE","MCS index","mcsIndex","0..31","31",SPEC,"A@_cuphyPuschUePrm:mcsIndex","")
ul(P,"UE","Modulation order Qm","qamModOrder","2/4/6/8; 1 (pi/2-BPSK) also allowed with transform precoding","8",DOC,"A@_cuphyPuschUePrm:qamModOrder","")
ul(P,"UE","Target code rate","targetCodeRate","codeRate*1024*10","-",DOC,"A@_cuphyPuschUePrm:targetCodeRate","")
ul(P,"UE","Redundancy version","rv","0..3","3",SPEC,"A@_cuphyPuschUePrm:rv","")
ul(P,"UE","HARQ process ID","harqProcessId","0..15","15",DOC,"A@_cuphyPuschUePrm:harqProcessId","")
ul(P,"UE","New data indicator","ndi","0 / 1","1",DOC,"A@_cuphyPuschUePrm:ndi","")
ul(P,"UE","Layers per UE","nUeLayers","1..4 (per 3GPP); <= 8 per UE group","4",SPEC,"A@_cuphyPuschUePrm:nUeLayers","")
ul(P,"UE","RNTI / data scrambling ID","rnti / dataScramId","0..65535 / 0..1023","-",SPEC,"A@_cuphyPuschUePrm:dataScramId","")
ul(P,"UE","LBRM","i_lbrm / maxLayers / maxQm / n_PRB_LBRM","as for PDSCH","-",DOC,"A@_cuphyPuschUePrm:i_lbrm","")
ul(P,"UE","Per-UE LDPC max iterations","ldpcMaxNumItrPerUe","uint8","-",DOC,"A@_cuphyPuschUePrm:ldpcMaxNumItrPerUe","Used when ldpcMaxNumItrAlgo = per UE.")
ul(P,"Transport block","Rate-matched bits per CB (Er)","PUSCH_MAX_ER_PER_CB_BITS","<= 256000","256000",EN,"H:PUSCH_MAX_ER_PER_CB_BITS","Exceeding it sets CUPHY_PUSCH_STATUS_UNSUPPORTED_MAX_ER_PER_CB.")
ul(P,"Transport block","Decoded CB size","MAX_DECODED_CODE_BLOCK_BIT_SIZE","<= 8448 bits","8448",CT_,"H:MAX_DECODED_CODE_BLOCK_BIT_SIZE","22 x 384.")
P = "UCI on PUSCH"
ul(P,"UCI","UCI-on-PUSCH UEs per cell group","CUPHY_MAX_N_UCI_ON_PUSCH","1..192","192 (256 with 64C)",CT_,"H:CUPHY_MAX_N_UCI_ON_PUSCH","")
ul(P,"UCI","HARQ-ACK / CSI-P1 bits","nBitsHarq / nBitsCsi1","uint16","-",DOC,"A@_cuphyUciOnPusch:nBitsHarq","Decoder chosen by size (see Algorithms).")
ul(P,"UCI","Beta offsets","betaOffsetHarqAck / betaOffsetCsi1 / betaOffsetCsi2","0..15 / 0..18 / 0..18","-",SPEC,"A@_cuphyUciOnPusch:betaOffsetHarqAck","")
ul(P,"UCI","Alpha scaling","alphaScaling","0..3","3",SPEC,"A@_cuphyUciOnPusch:alphaScaling","")
ul(P,"UCI","CSI-P2 reports","nCsi2Reports","0..100 (API), <= 16 per UE (buffers)","16",CT_,"H:CUPHY_MAX_N_CSI2_REPORTS_PER_UE","")
ul(P,"UCI","CSI-P2 size maps per cell","CUPHY_MAX_NUM_CSI2_SIZE_MAPS_PER_CELL","1..16","16",CT_,"H:CUPHY_MAX_NUM_CSI2_SIZE_MAPS_PER_CELL","")
ul(P,"UCI","DTX threshold","DTXthreshold","float (default -100)","-",DOC,"H:CUPHY_DEFAULT_EXT_DTX_THRESHOLD","")
ul(P,"UCI","CSI-RS ports (cell static, for CSI-P2 size)","nCsirsPorts","2, 4, 8, 12, 16, 24, 32","32",DOC,"H@_cuphyPuschCellStatPrm:nCsirsPorts","")
ul(P,"UCI","Codebook type","codebookType","0: Type1 single panel, 1: Type1 multi panel, 2: Type2, 3: Type2 port selection","3",DOC,"H@_cuphyPuschCellStatPrm:codebookType","")

P = "PUCCH"
ul(P,"Cell group","BBU layers","MAX_N_BBU_LAYERS_PUCCH_SUPPORTED","1..16","16",CT_,"H:MAX_N_BBU_LAYERS_PUCCH_SUPPORTED","")
ul(P,"Format 0","UCI groups","CUPHY_PUCCH_F0_MAX_GRPS","10 per cell","10 x cells",CT_,"H:CUPHY_PUCCH_F0_MAX_GRPS","")
ul(P,"Format 0","UCIs per group","CUPHY_PUCCH_F0_MAX_UCI_PER_GRP","1..12","12",CT_,"H:CUPHY_PUCCH_F0_MAX_UCI_PER_GRP","")
ul(P,"Format 1","UCI groups","CUPHY_PUCCH_F1_MAX_GRPS","24 per cell","24 x cells",CT_,"H:CUPHY_PUCCH_F1_MAX_GRPS","")
ul(P,"Format 1","UCIs per group","CUPHY_PUCCH_F1_MAX_UCI_PER_GRP","1..45","45",CT_,"H:CUPHY_PUCCH_F1_MAX_UCI_PER_GRP","")
ul(P,"Format 2","UCIs","CUPHY_PUCCH_F2_MAX_UCI","24 per cell","24 x cells",CT_,"H:CUPHY_PUCCH_F2_MAX_UCI","")
ul(P,"Format 2","PRBs","CUPHY_PUCCH_F2_MAX_PRBS","1..16","16",CT_,"H:CUPHY_PUCCH_F2_MAX_PRBS","")
ul(P,"Format 3","UCIs","CUPHY_PUCCH_F3_MAX_UCI","24 per cell","24 x cells",CT_,"H:CUPHY_PUCCH_F3_MAX_UCI","")
ul(P,"Format 3","PRBs","CUPHY_PUCCH_F3_MAX_PRBS","1..16","16",CT_,"H:CUPHY_PUCCH_F3_MAX_PRBS","")
ul(P,"Format 4","UCIs","nF4Ucis","API field present","-",DOC,"A:nF4Ucis","No dedicated Format 4 receiver directory in cuPHY/src/cuphy.")
ul(P,"Per UCI","Format type","formatType","0..4","4",SPEC,"H@_cuphyPucchUciPrm:formatType","")
ul(P,"Per UCI","BWP start/size","bwpStart","1..275","275",DOC,"H@_cuphyPucchUciPrm:bwpStart","")
ul(P,"Per UCI","Start symbol","startSym","0..13","13",SPEC,"H@_cuphyPucchUciPrm:startSym","")
ul(P,"Per UCI","Symbols","nSym","F0/F2: 1..2; F1/F3/F4: 4..14","14",SPEC,"H@_cuphyPucchUciPrm:nSym","")
ul(P,"Per UCI","Initial cyclic shift","initialCyclicShift","0..11","11",SPEC,"H@_cuphyPucchUciPrm:initialCyclicShift","")
ul(P,"Per UCI","Time-domain OCC index","timeDomainOccIdx","0..6","6",SPEC,"H@_cuphyPucchUciPrm:timeDomainOccIdx","F1.")
ul(P,"Per UCI","Hopping flags","freqHopFlag / groupHopFlag / sequenceHopFlag","0 / 1","1",SPEC,"H@_cuphyPucchUciPrm:freqHopFlag","")
ul(P,"Per UCI","pi/2-BPSK, additional DMRS","pi2Bpsk / AddDmrsFlag","0 / 1","1",SPEC,"H@_cuphyPucchUciPrm:pi2Bpsk","F3/F4.")
ul(P,"Per UCI","Max code rate","maxCodeRate","0..7","7",SPEC,"H@_cuphyPucchUciPrm:maxCodeRate","")
ul(P,"Per UCI","HARQ bits (F0/F1)","bitLenHarq","0..2","2",SPEC,"H@_cuphyPucchUciPrm:bitLenHarq","")
ul(P,"Per UCI","Scrambling IDs","dataScramblingId / DmrsScramblingId","0..1023 / 0..65535","-",SPEC,"H@_cuphyPucchUciPrm:dataScramblingId","")
ul(P,"Per UCI","UCI part-2 reports","numPart2s","0..100 (currently assumes <= 1)","1",DOC,"H:numPart2s","")
ul(P,"Cell","Hopping ID","pucchHoppingId","0..1023","1023",SPEC,"A@_cuphyPucchCellDynPrm:pucchHoppingId","")
ul(P,"Static","Polar list size","polarDcdrListSz","1..8 (PUCCH default 8)","8",DOC,"A@_cuphyPucchStatPrms:polarDcdrListSz","")
ul(P,"Static","UCI output mode","uciOutputMode","0: single buffer, 1: split HARQ/SR/CSI-P1","1",DOC,"A@_cuphyPucchStatPrms:uciOutputMode","")

P = "PRACH"
ul(P,"Cell static","Preamble formats","configurationIndex -> format","Format 0 (long) and B4 (short) only","-",EN,"PR:PreambleFormat::B4","Other formats rejected by validateStaticParams().")
ul(P,"Cell static","Configuration index","configurationIndex","0..255","255",DOC,"A@_cuphyPrachCellStatPrms:configurationIndex","")
ul(P,"Cell static","Numerology mu","mu","0 / 1","1",EN,"PR:cellPrms->mu","")
ul(P,"Cell static","Frequency range","FR","1 / 2","2",DOC,"A@_cuphyPrachCellStatPrms:FR","")
ul(P,"Cell static","Duplex","duplex","0: FDD, 1: TDD","1",DOC,"A@_cuphyPrachCellStatPrms:duplex","")
ul(P,"Cell static","Restricted set","restrictedSet","0 only (unrestricted)","0",EN,"PR:restrictedSet","")
ul(P,"Cell static","FDM occasions per cell","nFdmOccasions","1..8","8",DOC,"A@_cuphyPrachCellStatPrms:nFdmOccasions","")
ul(P,"Cell static","Antennas","N_ant","1..64","64",CT_,"H:MAX_N_ANTENNAS_SUPPORTED","")
ul(P,"Occasion","Root sequence index","prachRootSequenceIndex","0..137 (short) / 0..837 (long)","837",EN,"PR:prachRootSequenceIndex","Validated < L_RA - 1.")
ul(P,"Occasion","Zero-correlation zone config","prachZeroCorrConf","0..15","15",EN,"PR:prachZeroCorrConf","")
ul(P,"Occasion","Detection threshold override","force_thr0","0 = cuPHY default, > 0 = override","-",DOC,"A:force_thr0","")
ul(P,"Format 0","Sequence length / SCS / FFT","L_RA, delta_f_RA, Nfft","839 / 1.25 kHz / 1024; N_rep 1","-",CT_,"PR:table_NCS_1p25k","")
ul(P,"Format B4","Sequence length / SCS / FFT","L_RA, delta_f_RA, Nfft","139 / 15*2^mu kHz / 256; N_rep 12","-",CT_,"PR:table_NCS_15kplus","")
ul(P,"Output","Preambles per occasion","CUPHY_PRACH_RX_NUM_PREAMBLE","0..64","64",CT_,"A:CUPHY_PRACH_RX_NUM_PREAMBLE","Outputs index, delay, power, RSSI, per-antenna RSSI, interference.")

P = "SRS (Rx)"
ul(P,"Cell group","SRS UEs","nSrsUes / CUPHY_SRS_MAX_N_USERS","1..512","512",CT_,"H:CUPHY_SRS_MAX_N_USERS","")
ul(P,"Cell group","SRS cells","MAX_N_SRS_CELL","1..24","24",CT_,"SRS:MAX_N_SRS_CELL","")
ul(P,"Cell","SRS symbols per slot","nSrsSym / MAX_SRS_SYMBOLS_PER_SLOT","1..4","4",CT_,"H:MAX_SRS_SYMBOLS_PER_SLOT","")
ul(P,"Cell","Full-band SRS antenna ports per slot per cell","CUPHY_SRS_MAX_FULL_BAND_SRS_ANT_PORTS_SLOT_PER_CELL","1..16","16",CT_,"H:CUPHY_SRS_MAX_FULL_BAND_SRS_ANT_PORTS_SLOT_PER_CELL","")
ul(P,"Cell","Full-band channel estimates per TTI","CUPHY_SRS_MAX_FULL_BAND_CHEST_PER_TTI","1..8","8",CT_,"H:CUPHY_SRS_MAX_FULL_BAND_CHEST_PER_TTI","")
ul(P,"Cell","PRGs","CUPHY_SRS_MAX_N_PRGS_SUPPORTED","1..273","273",CT_,"H:CUPHY_SRS_MAX_N_PRGS_SUPPORTED","")
ul(P,"Static","Channel estimation algorithm","chEstAlgo","0: MMSE, 1: RKHS","1",DOC,"H:SRS_CH_EST_ALGO_TYPE_RKHS","")
ul(P,"Static","Normalization to L2","chEstToL2NormalizationAlgo","0: constant scaling, 1: peak normalization","1",DOC,"A:chEstToL2NormalizationAlgo","")
ul(P,"Static","Delay offset correction","enableDelayOffsetCorrection","0 / 1","1",DOC,"A:enableDelayOffsetCorrection","")
ul(P,"UE","Antenna ports","nAntPorts","1, 2, 4","4",DOC,"H@_cuphyUeSrsPrm:nAntPorts","")
ul(P,"UE","Symbols","nSyms","1, 2, 4","4",DOC,"H@_cuphyUeSrsPrm:nSyms","")
ul(P,"UE","Repetitions","nRepetitions","1, 2, 4","4",DOC,"H@_cuphyUeSrsPrm:nRepetitions","")
ul(P,"UE","Comb size","combSize","2, 4","4",DOC,"H@_cuphyUeSrsPrm:combSize","")
ul(P,"UE","Comb offset","combOffset","0..3","3",DOC,"H@_cuphyUeSrsPrm:combOffset","")
ul(P,"UE","Start symbol","startSym","0..13","13",DOC,"H@_cuphyUeSrsPrm:startSym","")
ul(P,"UE","Sequence ID","sequenceId","0..1023","1023",DOC,"H@_cuphyUeSrsPrm:sequenceId","")
ul(P,"UE","Bandwidth config index (C_SRS)","configIdx","0..63","63",DOC,"H@_cuphyUeSrsPrm:configIdx","")
ul(P,"UE","Bandwidth index (B_SRS)","bandwidthIdx","0..3","3",DOC,"H@_cuphyUeSrsPrm:bandwidthIdx","")
ul(P,"UE","Cyclic shift","cyclicShift","0..11","11",DOC,"H@_cuphyUeSrsPrm:cyclicShift","")
ul(P,"UE","Frequency position","frequencyPosition","0..67","67",DOC,"H@_cuphyUeSrsPrm:frequencyPosition","")
ul(P,"UE","Frequency shift","frequencyShift","0..268","268",DOC,"H@_cuphyUeSrsPrm:frequencyShift","")
ul(P,"UE","Frequency hopping","frequencyHopping","0..3","3",DOC,"H@_cuphyUeSrsPrm:frequencyHopping","")
ul(P,"UE","Resource type","resourceType","0: aperiodic, 1: semi-persistent, 2: periodic","2",DOC,"H@_cuphyUeSrsPrm:resourceType","")
ul(P,"UE","Periodicity T_SRS (slots)","Tsrs","0,2,3,5,8,10,16,20,32,40,64,80,160,320,640,1280,2560","2560",DOC,"H@_cuphyUeSrsPrm:Tsrs","")
ul(P,"UE","Slot offset","Toffset","0..2569","2569",DOC,"H@_cuphyUeSrsPrm:Toffset","")
ul(P,"UE","Group/sequence hopping","groupOrSequenceHopping","0: none, 1: group, 2: sequence","2",DOC,"H@_cuphyUeSrsPrm:groupOrSequenceHopping","")
ul(P,"UE","Frequency hops","MAX_N_HOPS","1..4","4",CT_,"SRS:MAX_N_HOPS","")

# ============================ ALGORITHMS ============================
ALG = []
def al(*a): ALG.append(a)
# (direction, channel, processing block, algorithm / options, selector / key params, limits, src, dir)
al("DL","PDSCH","TB/CB CRC attach","CRC24A (TB), CRC24B (CB)","read_TB_CRC","-","A:read_TB_CRC","cuPHY/src/cuphy/crc")
al("DL","PDSCH","LDPC encoding","5G NR QC-LDPC, base graph 1 / 2","BG, Zc","Zc <= 384; BG1 22 info / 46 parity columns; BG2 10 info / 42 parity; encoded CB <= 25344 bits","H:CUPHY_LDPC_MAX_LIFTING_SIZE","cuPHY/src/cuphy/ldpc")
al("DL","PDSCH","Rate matching","Circular buffer + bit interleaving; LBRM (38.212 5.4.2)","rv, n_PRB_LBRM, maxLayers, maxQm","Er <= 256000 bits per CB","H:PDSCH_MAX_ER_PER_CB_BITS","cuPHY/src/cuphy/dl_rate_matching")
al("DL","PDSCH","Scrambling","Gold sequence (c_init from RNTI, dataScramId)","rnti, dataScramId","-","","cuPHY/src/cuphy/Gold_sequence")
al("DL","PDSCH","Modulation mapping","QPSK, 16QAM, 64QAM, 256QAM","qamModOrder","Qm <= 8","H:CUPHY_QAM_256","cuPHY/src/cuphy/modulation_mapper")
al("DL","PDSCH","Layer mapping + precoding","Per-UE precoder matrix (fp16)","enablePrcdBf, pmwPrmIdx","<= 4 layers/TB, <= 8 layers/UE, <= 32 ports","H:MAX_DL_PORTS","cuPHY/src/cuphy/pdsch_dmrs")
al("DL","PDSCH","DMRS generation","Type-1 DMRS, fOCC/tOCC","dmrsPortBmsk, scid, dmrsScrmId","<= 4 DMRS symbols","H:MAX_N_DMRSSYMS_SUPPORTED","cuPHY/src/cuphy/pdsch_dmrs")
al("DL","PDSCH","CSI-RS RE rate matching","Skip REs occupied by co-scheduled CSI-RS","nCsiRsPrms","<= 32 CSI-RS params per cell","H:CUPHY_CSIRS_MAX_NUM_PARAMS","cuPHY/src/cuphy/csirs")
al("DL","PDCCH","DCI encoding","CRC24C attach + RNTI mask, Polar encoding (38.212 7.3)","Npayload, rntiCrc","A 12..140, K <= 164, N <= 512","H:CUPHY_POLAR_ENC_MAX_INFO_BITS","cuPHY/src/cuphy/polar_encoder")
al("DL","PDCCH","Rate matching + scrambling + QPSK","Polar rate matching, Gold scrambling","aggr_level","Tx bits <= 1728 per DCI (AL 16)","H:CUPHY_PDCCH_MAX_TX_BITS_PER_DCI","cuPHY/src/cuphy/pdcch")
al("DL","PDCCH","CCE-to-REG mapping","Interleaved / non-interleaved","interleaved, bundle_size, interleaver_size, shift_index","-","A@_cuphyPdcchCoresetDynPrm:interleaved","cuPHY/src/cuphy/pdcch")
al("DL","SSB","PSS/SSS generation","m-sequences (length 127)","NID","-","H:CUPHY_SSB_N_SS_SEQ_BITS","cuPHY/src/cuphy/ss")
al("DL","SSB","PBCH encoding","Payload gen + CRC24C + Polar (N=512) + rate match (864) + scrambling","MIB, SFN, nHF, blockIndex","<= 3 SSB/cell/slot","H:CUPHY_SSB_N_PBCH_POLAR_ENCODED_BITS","cuPHY/src/cuphy/ss")
al("DL","CSI-RS","CSI-RS generation / mapping","38.211 Table 7.4.1.5.3-1 rows 1..18; CDM none/2/4/8","row, cdmType, freqDensity","<= 32 ports","H:CUPHY_CSIRS_MAX_ANTENNA_PORTS","cuPHY/src/cuphy/csirs")
al("DL/UL","BFW","Beamforming weight computation","Regularized zero-forcing from SRS channel estimates; optional 9-bit BFP compression","lambda, compressBitwidth, beta, bfwPowerNormAlg_selector","<= 72 UE groups, <= 32 layers/UE group, <= 137 PRB groups, <= 64 antennas","H:CUPHY_BFW_COEF_COMP_N_MAX_USER_GRPS","cuPHY/src/cuphy/bfc")
al("UL","PUSCH","Channel estimation","0 legacy MMSE (freq-interp filters); 1 multi-stage MMSE + delay estimation; 2 RKHS; 3 LS only","chEstAlgo, enablePerPrgChEst, enablePuschTdi","<= 16 het configs (legacy), <= 4 (multi-stage); <= 4 time estimates","H:PUSCH_CH_EST_ALGO_TYPE_LEGACY_MMSE","cuPHY/src/cuphy/ch_est")
al("UL","PUSCH","Noise / interference estimation","Per-PRB noise-plus-interference covariance from DMRS","-","1 het config","H:CUPHY_PUSCH_RX_NOISE_INTF_EST_N_MAX_HET_CFGS","cuPHY/src/cuphy/pusch_noise_intf_est")
al("UL","PUSCH","CFO / TA estimation and correction","Phase-rotation-based CFO; weighted-average CFO; timing offset","enableCfoCorrection, enableWeightedAverageCfo, enableToEstimation, foForgetCoeff","-","H:CUPHY_PUSCH_RX_CFO_EST_N_MAX_HET_CFGS","cuPHY/src/cuphy/cfo_ta_est")
al("UL","PUSCH","Equalization","0 RZF; 1 noise-diagonal MMSE; 2 MMSE-IRC; 3 MMSE-IRC + RBLW shrinkage; 4 MMSE-IRC + OAS shrinkage","eqCoeffAlgo","<= 8 het configs; <= 8 layers per UE group","H:PUSCH_EQ_ALGO_TYPE_RZF","cuPHY/src/cuphy/channel_eq")
al("UL","PUSCH","DFT-s-OFDM de-precoding","IDFT after equalization (transform precoding)","enableDftSOfdm, enableTfPrcd","-","A:enableDftSOfdm","cuPHY/src/cuphy/channel_eq")
al("UL","PUSCH","Soft demapping","LLR computation for pi/2-BPSK..256QAM; early-HARQ symbol subset","qamModOrder","Qm <= 8","H:CUPHY_PUSCH_RX_SOFT_DEMAPPER_FULL_SLOT_SYMBOL_BITMASK","cuPHY/src/cuphy/soft_demapper")
al("UL","PUSCH","Measurements","RSSI, RSRP, SINR","enableRssiMeasurement, enableSinrMeasurement, rssiSymLocBmsk","1 het config each","H:CUPHY_PUSCH_RX_RSSI_N_MAX_HET_CFGS","cuPHY/src/cuphy/pusch_rssi")
al("UL","PUSCH","De-rate-matching + HARQ combining","Inverse of 38.212 rate matching, LBRM","rv, ndi, harqProcessId","Er <= 256000 bits per CB","H:PUSCH_MAX_ER_PER_CB_BITS","cuPHY/src/cuphy/rate_matching")
al("UL","PUSCH","LDPC decoding","Layered min-sum decoder; TB or CB mode; early termination; fp16 option","ldpcAlgoIndex, ldpcFlags, ldpcEarlyTermination, ldpcUseHalf, ldpcMaxNumItrAlgo (fixed/LUT/per-UE), ldpcClampValue","Zc <= 384; decoded CB <= 8448 bits","H:CUPHY_LDPC_DECODE_EARLY_TERM","cuPHY/src/cuphy/ldpc")
al("UL","PUSCH","CRC check","TB/CB CRC verification","-","-","","cuPHY/src/cuphy/crc")
al("UL","UCI (PUSCH/PUCCH)","Small-block decoding","Simplex decoder (1..2 bits); Reed-Muller decoder (3..11 bits)","nInfoBits","Simplex <= 2 bits, RM <= 11 bits","H:CUPHY_N_MAX_UCI_BITS_RM","cuPHY/src/cuphy/simplex_decoder")
al("UL","UCI (PUSCH/PUCCH)","Polar decoding","CRC-aided successive-cancellation list decoder; segmentation/de-rate-match/de-interleave","polarDcdrListSz","List <= 8; N <= 1024","H:CUPHY_POLAR_DECODER_MAX_BITS","cuPHY/src/cuphy/polar_decoder")
al("UL","UCI (PUSCH/PUCCH)","DTX detection","Threshold-based DTX for HARQ/CSI-1/CSI-2","DTXthreshold","Default -100","H:CUPHY_DEFAULT_EXT_DTX_THRESHOLD","cuPHY/src/cuphy/uci_on_pusch")
al("UL","PUCCH","Format 0 receiver","Sequence-correlation detection (HARQ/SR)","initialCyclicShift","<= 12 UCIs/group, 10 groups/cell","H:CUPHY_PUCCH_F0_MAX_UCI_PER_GRP","cuPHY/src/cuphy/pucch_F0_receiver")
al("UL","PUCCH","Format 1 receiver","Channel estimation + OCC despreading + BPSK/QPSK detection","timeDomainOccIdx","<= 45 UCIs/group, 24 groups/cell","H:CUPHY_PUCCH_F1_MAX_UCI_PER_GRP","cuPHY/src/cuphy/pucch_F1_receiver")
al("UL","PUCCH","Format 2 front end","DMRS channel estimation, equalization, LLR","prbSize","<= 16 PRBs","H:CUPHY_PUCCH_F2_MAX_PRBS","cuPHY/src/cuphy/pucch_F2_front_end")
al("UL","PUCCH","Format 3 front end","DFT-s-OFDM demod + LLR; CSI-P2 control","prbSize, pi2Bpsk, AddDmrsFlag","<= 16 PRBs","H:CUPHY_PUCCH_F3_MAX_PRBS","cuPHY/src/cuphy/pucch_F3_front_end")
al("UL","PRACH","Preamble detection","Zadoff-Chu correlation via FFT (cuFFTDx), non-coherent combining, peak search, threshold","prachRootSequenceIndex, prachZeroCorrConf, force_thr0","FFT 256 / 1024; formats 0 & B4; <= 64 preambles","PRK:PRACH_SUPPORTED_FFT_SIZES","cuPHY/src/cuphy/prach_receiver")
al("UL","SRS","Channel estimation","0 MMSE; 1 RKHS","chEstAlgo, chEstToL2NormalizationAlgo, enableDelayOffsetCorrection","<= 512 UEs, <= 4 ports/UE, <= 4 symbols","H:SRS_CH_EST_ALGO_TYPE_MMSE","cuPHY/src/cuphy/srs_chEst")
al("UE-side/test","CSI-RS Rx","CSI-RS channel estimation","Per-UE CSI-RS channel estimate","nRxAnt","<= 1024 UEs","H:CUPHY_CSIRS_MAX_NUM_UES","cuPHY/src/cuphy/csirs_rx")
al("UE-side/test","SRS Tx","SRS generation","Low-PAPR sequence generation","same as SRS UE params","-","","cuPHY/src/cuphy/srstx")

# ============================ GLOBAL ============================
GL = [
 ("System","Cells per slot (DL and UL)","MAX_CELLS_PER_SLOT","20","40",CT_,"H:MAX_CELLS_PER_SLOT"),
 ("System","Numerology mu (cell)","cuphyCellStatPrm_t.mu","0..3","-",DOC,"A@_cuphyCellStatPrm:mu"),
 ("System","OFDM symbols per slot","OFDM_SYMBOLS_PER_SLOT","14","14",CT_,"H:OFDM_SYMBOLS_PER_SLOT"),
 ("System","Subcarriers per PRB","CUPHY_N_TONES_PER_PRB","12","12",CT_,"H:CUPHY_N_TONES_PER_PRB"),
 ("System","PRBs per carrier","MAX_N_PRBS_SUPPORTED","273","273",CT_,"H:MAX_N_PRBS_SUPPORTED"),
 ("System","Antennas","MAX_N_ANTENNAS_SUPPORTED","64","64",CT_,"H:MAX_N_ANTENNAS_SUPPORTED"),
 ("System","Carriers","MAX_N_CARRIERS_SUPPORTED","10","10",CT_,"H:MAX_N_CARRIERS_SUPPORTED"),
 ("System","BBU layers (all)","MAX_N_BBU_LAYERS_SUPPORTED","32","32",CT_,"H:MAX_N_BBU_LAYERS_SUPPORTED"),
 ("System","TBs per cell group","MAX_N_TBS_PER_CELL_GROUP_SUPPORTED","192","256",CT_,"H:MAX_N_TBS_PER_CELL_GROUP_SUPPORTED"),
 ("System","UE groups","MAX_N_USER_GROUPS_SUPPORTED","192","256",CT_,"H:MAX_N_USER_GROUPS_SUPPORTED"),
 ("System","CBs per TB","MAX_N_CBS_PER_TB_SUPPORTED","152","152",CT_,"H:MAX_N_CBS_PER_TB_SUPPORTED"),
 ("LDPC","Max lifting size Zc","CUPHY_LDPC_MAX_LIFTING_SIZE","384","384",CT_,"H:CUPHY_LDPC_MAX_LIFTING_SIZE"),
 ("LDPC","BG1 info / parity / variable nodes","CUPHY_LDPC_BG1_INFO_NODES ...","22 / 46 / 68","-",CT_,"H:CUPHY_LDPC_BG1_INFO_NODES"),
 ("LDPC","BG2 info / parity / variable nodes","CUPHY_LDPC_MAX_BG2_INFO_NODES ...","10 / 42 / 52","-",CT_,"H:CUPHY_LDPC_MAX_BG2_INFO_NODES"),
 ("LDPC","Min parity nodes","CUPHY_LDPC_MIN_PARITY_NODES","4","-",CT_,"H:CUPHY_LDPC_MIN_PARITY_NODES"),
 ("LDPC","Max encoded CB bits","MAX_ENCODED_CODE_BLOCK_BIT_SIZE","25344","-",CT_,"H:MAX_ENCODED_CODE_BLOCK_BIT_SIZE"),
 ("LDPC","Max decoded CB bits","MAX_DECODED_CODE_BLOCK_BIT_SIZE","8448","-",CT_,"H:MAX_DECODED_CODE_BLOCK_BIT_SIZE"),
 ("LDPC","TBs per decode descriptor","CUPHY_LDPC_DECODE_DESC_MAX_TB","32","-",CT_,"H:CUPHY_LDPC_DECODE_DESC_MAX_TB"),
 ("Polar","Encoder max info / coded / tx bits","CUPHY_POLAR_ENC_MAX_*","164 / 512 / 8192","-",CT_,"H:CUPHY_POLAR_ENC_MAX_INFO_BITS"),
 ("Polar","Decoder max N / list size","CUPHY_POLAR_DECODER_MAX_BITS / _LIST_SIZE","1024 / 8","-",CT_,"H:CUPHY_POLAR_DECODER_MAX_BITS"),
 ("Modulation","QAM orders (bits/symbol)","CUPHY_QAM_2..256","1, 2, 4, 6, 8","-",CT_,"H:CUPHY_QAM_2"),
]

# ============================ WRITE ============================
wb = Workbook()
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF")

def sheet(ws, headers, rows, widths, table_name):
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"
    ref_ = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    t = Table(displayName=table_name, ref=ref_)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)

# README
ws = wb.active; ws.title = "README"
readme = [
 ["cuPHY 5G NR channels, signals, algorithms and parameter limits"],
 [""],
 ["Scope", "Parameters exposed by the cuPHY L1 API (cuPHY/src/cuphy/cuphy_api.h, cuphy.h) and the checks in cuPHY/src/cuphy_channels/*.cpp."],
 ["Sheets", "Global_Limits: system-wide compile-time maxima. Downlink / Uplink: per channel/signal parameters. Algorithms: processing blocks and selectable algorithms."],
 ["Basis column", f"'{CT_}': #define/constexpr that sizes buffers. '{EN}': value checked at create/setup time; out-of-range values are rejected. '{DOC}': range stated in the API header comment, not necessarily checked. '{SPEC}': the 3GPP/SCF-FAPI range for a field whose cuPHY comment gives no range."],
 ["64C build", "Several limits are larger when built with ENABLE_64C (64-cell configuration); both values are shown where they differ."],
 ["Source column", "path:line into the aerial-cuda-accelerated-ran repository, resolved when this workbook was generated."],
 ["Caveat", "cuPHY-CP (cuphydriver/FAPI adapter) and deployment configs may impose tighter limits than cuPHY itself."],
]
for r in readme: ws.append(r)
ws["A1"].font = Font(bold=True, size=14)
for r in range(3, len(readme) + 1):
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.cell(row=r, column=2).alignment = WRAP
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 120

H_CH = ["Channel / Signal", "Group", "Parameter", "cuPHY field / macro", "Allowed values / range", "Max value", "Basis", "Source", "Notes"]
W_CH = [20, 16, 34, 40, 44, 16, 22, 48, 44]

ws = wb.create_sheet("Global_Limits")
sheet(ws, ["Area", "Parameter", "cuPHY macro", "Default build", "ENABLE_64C build", "Basis", "Source"],
      [(a, b, c, d, e, f, ref(g)) for a, b, c, d, e, f, g in GL], [14, 36, 44, 18, 18, 20, 48], "GlobalLimits")

ws = wb.create_sheet("Downlink")
sheet(ws, H_CH, [(*r[:7], ref(r[7]), r[8]) for r in DL], W_CH, "Downlink")

ws = wb.create_sheet("Uplink")
sheet(ws, H_CH, [(*r[:7], ref(r[7]), r[8]) for r in UL], W_CH, "Uplink")

ws = wb.create_sheet("Algorithms")
sheet(ws, ["Direction", "Channel", "Processing block", "Algorithm / options", "Selector / key parameters", "Limits", "Source", "Implementation dir"],
      [(a, b, c, d, e, f, ref(g), h) for a, b, c, d, e, f, g, h in ALG], [12, 18, 30, 52, 42, 40, 48, 36], "Algorithms")

wb.save(OUT)
print("wrote", OUT, "DL rows", len(DL), "UL rows", len(UL), "ALG rows", len(ALG), "GL rows", len(GL))
